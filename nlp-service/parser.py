import re
import pypdf
import docx

# Predefined list of standard tech skills for basic local matching
COMMON_SKILLS = [
    # Programming Languages
    "python", "javascript", "typescript", "java", "c++", "c#", "go", "golang", "ruby", "php", "swift", "kotlin", "rust", "scala", "sql", "html", "css", "r",
    # Frameworks & Libraries
    "react", "angular", "vue", "next.js", "nextjs", "express", "express.js", "node.js", "nodejs", "django", "flask", "fastapi", "spring boot", "laravel", "rails", "asp.net", "nest.js", "react native", "flutter", "jquery", "bootstrap", "tailwind", "redux", "graphql",
    # Databases & Caching
    "mongodb", "postgresql", "mysql", "sqlite", "redis", "elasticsearch", "cassandra", "dynamodb", "mariadb", "firebase", "firestore", "oracle",
    # Cloud, DevOps & Tools
    "aws", "amazon web services", "azure", "gcp", "google cloud", "docker", "kubernetes", "git", "github", "gitlab", "jenkins", "terraform", "ansible", "ci/cd", "nginx", "apache", "linux", "bash", "firebase",
    # AI/ML/Data Science
    "machine learning", "deep learning", "nlp", "natural language processing", "computer vision", "tensorflow", "pytorch", "keras", "scikit-learn", "numpy", "pandas", "spark", "hadoop", "tableau", "power bi", "openai", "llm", "gemini",
    # Architecture & Concepts
    "rest api", "microservices", "system design", "agile", "scrum", "oop", "dsa", "data structures", "algorithms"
]

def extract_text_from_pdf(stream):
    """Extracts text from PDF file stream using pypdf."""
    text = ""
    try:
        reader = pypdf.PdfReader(stream)
        for page in reader.pages:
            text += (page.extract_text() or "") + "\n"
    except Exception as e:
        print(f"Error reading PDF: {e}")
    return text


def extract_text_from_docx(stream):
    """Extracts text from DOCX file stream using python-docx."""
    text = ""
    try:
        doc = docx.Document(stream)
        for para in doc.paragraphs:
            text += para.text + "\n"
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    text += cell.text + " "
                text += "\n"
    except Exception as e:
        print(f"Error reading DOCX: {e}")
    return text

def parse_contact_info(text):
    """Extracts email, phone number, and standard profiles using regex."""
    email_pattern = r'[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+'
    phone_pattern = r'(?:\+?\d{1,3}[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}'
    github_pattern = r'github\.com/[a-zA-Z0-9_-]+'
    linkedin_pattern = r'linkedin\.com/in/[a-zA-Z0-9_-]+'

    email = re.search(email_pattern, text)
    phone = re.search(phone_pattern, text)
    github = re.search(github_pattern, text)
    linkedin = re.search(linkedin_pattern, text)

    return {
        "email": email.group(0) if email else "",
        "phone": phone.group(0) if phone else "",
        "github": github.group(0) if github else "",
        "linkedin": linkedin.group(0) if linkedin else ""
    }

def extract_skills_locally(text):
    """Matches the resume text against a list of common skills (case-insensitive)."""
    text_lower = text.lower()
    matched_skills = []
    
    # We look for word boundaries around the skills to avoid partial matches
    for skill in COMMON_SKILLS:
        pattern = r'\b' + re.escape(skill) + r'\b'
        if re.search(pattern, text_lower):
            # Normalize display name
            matched_skills.append(skill.upper() if skill in ["aws", "gcp", "dsa", "oop", "sql", "html", "css", "api", "nlp", "llm", "hr"] else skill.title())
            
    return list(set(matched_skills))

def parse_sections(text):
    """Finds content in key sections using keyword anchors."""
    sections = {
        "experience": "",
        "education": "",
        "projects": "",
        "skills": ""
    }
    
    lines = text.split('\n')
    current_section = None
    section_headers = {
        "experience": ["experience", "work history", "employment", "professional background", "internship", "internships"],
        "education": ["education", "academic background", "qualification", "qualifications", "schooling", "college", "university"],
        "projects": ["projects", "personal projects", "academic projects", "key projects"],
        "skills": ["skills", "technical skills", "core competencies", "technologies", "expertise"]
    }
    
    for line in lines:
        line_clean = line.strip().lower()
        if not line_clean:
            continue
        
        # Check if line looks like a header
        header_matched = False
        for sec, keywords in section_headers.items():
            for kw in keywords:
                # Header lines are usually short
                if len(line_clean) < 30 and (line_clean == kw or line_clean.startswith(kw + " ") or line_clean.endswith(" " + kw)):
                    current_section = sec
                    header_matched = True
                    break
            if header_matched:
                break
        
        if header_matched:
            continue
            
        if current_section:
            sections[current_section] += line + "\n"
            
    return sections
